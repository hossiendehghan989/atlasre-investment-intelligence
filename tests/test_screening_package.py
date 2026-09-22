from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from generate_committee_report import DEFAULT_RISK_SIMULATIONS, build_screening_package, screening_package_zip
from scripts.screen_deal import load_deal, run
from src.advanced_underwriting import monte_carlo_underwriting, risk_summary
from src.atlasre import DealInputs, validate_deal
from src.ic_workflow import DealCase, generate_ic_memo
from src.presentation import metric_help, number_or_na, percent_or_na
from src.reconciliation import build_reconciliation_workbook


def test_screening_package_contains_executive_sections_and_supporting_files():
    files = build_screening_package(DealInputs(10_000_000, 650_000, leverage=0.5), simulations=100)
    report = files["investment_committee_report.md"].decode()
    memo = files["investment_committee_memo.md"].decode()
    package_readme = files["README_ILLUSTRATIVE.md"].decode()
    assert report.startswith("# ILLUSTRATIVE")
    assert memo.startswith("# ILLUSTRATIVE")
    assert package_readme.startswith("# ILLUSTRATIVE")
    assert "Downside and governance flags" in report
    assert "Expected shortfall" in report
    assert "Model-run fingerprint" in report
    assert "review_required" not in report.lower()
    assert {
        "stress_cases.csv",
        "annual_debt_schedule.csv",
        "monthly_development_model.csv",
        "assumptions.csv",
        "lineage.json",
        "risk_summary.json",
        "lease_summary.csv",
        "lease_monthly_rollup.csv",
        "excel_reconciliation.xlsx",
    }.issubset(files)


def test_screening_package_zip_is_readable():
    package = screening_package_zip(DealInputs(10_000_000, 650_000, leverage=0.5), simulations=50)
    with ZipFile(BytesIO(package)) as archive:
        names = set(archive.namelist())
    assert "investment_committee_report.md" in names
    assert "investment_committee_memo.md" in names
    assert "README_ILLUSTRATIVE.md" in names
    assert "annual_debt_schedule.csv" in names
    assert "monthly_development_model.csv" in names
    assert "lease_summary.csv" in names
    assert "excel_reconciliation.xlsx" in names


def test_dashboard_and_cli_default_risk_numbers_match():
    deal = DealInputs(10_000_000, 650_000, hold_years=5, leverage=0.5)
    dashboard_risk = risk_summary(
        monte_carlo_underwriting(deal, simulations=DEFAULT_RISK_SIMULATIONS, seed=42),
        hurdle_rate=0.12,
    )
    cli_files = build_screening_package(deal, simulations=DEFAULT_RISK_SIMULATIONS)
    cli_risk = json.loads(cli_files["risk_summary.json"])
    assert dashboard_risk == cli_risk


def test_verified_without_evidence_is_review_required_throughout_package():
    files = build_screening_package(
        DealInputs(10_000_000, 650_000, leverage=0.5),
        source_status="VERIFIED",
        simulations=25,
    )
    report = files["investment_committee_report.md"].decode()
    assumptions = files["assumptions.csv"].decode()
    assert "| Source status | **REVIEW REQUIRED** |" in report
    assert "Source-backed screening input" not in assumptions
    assert "REVIEW REQUIRED" in assumptions


def test_verified_with_evidence_is_verified_throughout_package():
    files = build_screening_package(
        DealInputs(10_000_000, 650_000, leverage=0.5),
        source_status="VERIFIED",
        verified_by="Reviewer A",
        source_reference="data-room://deal-001/operating-statement.pdf",
        simulations=25,
    )
    report = files["investment_committee_report.md"].decode()
    assumptions = files["assumptions.csv"].decode()
    assert "| Source status | **VERIFIED** |" in report
    assert "Source-backed screening input" in assumptions
    assert "VERIFIED" in assumptions


# Dashboard and review-package regressions


ROOT = Path(__file__).resolve().parents[1]


def test_risk_summary_handles_no_debt_series_independently():
    from src.advanced_underwriting import monte_carlo_underwriting, risk_summary
    from src.atlasre import DealInputs

    summary = risk_summary(monte_carlo_underwriting(DealInputs(10_000_000, 650_000, leverage=0), simulations=20))
    assert summary["probability_dscr_below_125"] is None
    assert summary["p05_irr"] is not None


def test_dashboard_regressions_and_illustrative_notice():
    from src.atlasre import underwrite_deal

    for deal in (DealInputs(10_000_000, 650_000, leverage=0.0), DealInputs(500_000_000, 650_000, leverage=0.5)):
        validate_deal(deal)
        result = underwrite_deal(deal)
        assert "inf" not in number_or_na(result["minimum_dscr"]).lower()


@pytest.mark.parametrize(
    ("label", "target", "message"),
    [
        ("Purchase price ($)", 0, "purchase_price and annual_noi must be positive"),
        ("Purchase price ($)", 10**15, "purchase_price and annual_noi must not exceed"),
    ],
)
def test_dashboard_shows_clear_validation_error_without_traceback(label, target, message):
    deal = DealInputs(target, 650_000) if label == "Purchase price ($)" else DealInputs(10_000_000, 650_000)
    with pytest.raises(ValueError, match=message):
        validate_deal(deal)


def test_dashboard_does_not_substitute_unlevered_irr_for_nonconvergent_levered_irr():
    from src.atlasre import underwrite_deal

    result = underwrite_deal(DealInputs(10_000_000_000, 650_000, leverage=0.5))
    assert percent_or_na(result["levered_irr"]) == "N/A"
    assert percent_or_na(result["unlevered_irr"]) == "-73.71%"
    assert metric_help(result["levered_irr"], "IRR did not converge") == "IRR did not converge"


def test_generated_package_does_not_substitute_nonconvergent_irr():
    files = build_screening_package(DealInputs(10_000_000_000, 650_000, leverage=0.5), simulations=10)
    report = files["investment_committee_report.md"].decode()
    memo = files["investment_committee_memo.md"].decode()

    assert "| Levered IRR | N/A — IRR did not converge |" in report
    assert "| Unlevered IRR | -73.71% |" in report
    assert "| Levered IRR | N/A — IRR did not converge |" in memo


def test_template_validates_and_runs_end_to_end(tmp_path):
    template_path = ROOT / "docs/templates/deal_template.json"
    payload = json.loads(template_path.read_text())
    payload["simulations"] = 10
    input_path = tmp_path / "deal.json"
    input_path.write_text(json.dumps(payload))
    assert load_deal(input_path)["purchase_price"] == 10_000_000
    output = run(input_path, tmp_path / "review")
    assert (output / "investment_committee_memo.md").exists()
    assert (output / "assumptions.csv").exists()
    assert (output / "risk_summary.json").exists()
    assert (output / "annual_debt_schedule.csv").exists()
    assert (output / "excel_reconciliation.xlsx").exists()


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("purchase_price", None, "missing required fields"),
        ("annual_noi", float("nan"), "must be finite"),
        ("leverage", -0.1, "cannot be negative"),
        ("hold_years", 5.5, "must be an integer"),
        ("simulations", 0, "positive integer"),
    ],
)
def test_invalid_input_is_rejected_clearly(tmp_path, field, value, message):
    payload = json.loads((ROOT / "docs/templates/deal_template.json").read_text())
    if value is None:
        payload.pop(field)
    else:
        payload[field] = value
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(payload, allow_nan=True))
    with pytest.raises(ValueError, match=message):
        load_deal(path)


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("hold_years", 10**15, "must not exceed"),
        ("annual_noi_growth", 1e300, "no greater than 100%"),
        ("debt_rate", 1e15, "debt rate"),
        ("debt_amortization_years", 10**15, "must not exceed"),
        ("simulations", 10**15, "between 1 and"),
    ],
)
def test_hostile_numeric_magnitudes_are_rejected_before_package_generation(tmp_path, field, value, message):
    payload = json.loads((ROOT / "docs/templates/deal_template.json").read_text())
    payload[field] = value
    path = tmp_path / "hostile-numeric.json"
    path.write_text(json.dumps(payload))

    with pytest.raises(ValueError, match=message):
        run(path, tmp_path / "review")


def test_verified_source_gate_requires_both_evidence(tmp_path):
    payload = json.loads((ROOT / "docs/templates/deal_template.json").read_text())
    payload["simulations"] = 10
    payload["source_status"] = "VERIFIED"
    payload["verified_by"] = ""
    payload["source_reference"] = ""
    input_path = tmp_path / "deal.json"
    input_path.write_text(json.dumps(payload))
    report = (run(input_path, tmp_path / "review") / "investment_committee_report.md").read_text()
    assert "Source status | **REVIEW REQUIRED**" in report
    assert "Decision status" in report


def test_cli_verified_source_gate_passes_only_with_both_evidence(tmp_path):
    payload = json.loads((ROOT / "docs/templates/deal_template.json").read_text())
    payload.update(
        {
            "simulations": 10,
            "source_status": "VERIFIED",
            "verified_by": "Owner reviewer",
            "source_reference": "local://source.pdf",
        }
    )
    input_path = tmp_path / "verified.json"
    input_path.write_text(json.dumps(payload))
    report = (run(input_path, tmp_path / "verified-review") / "investment_committee_report.md").read_text()
    assert "Source status | **VERIFIED**" in report


# Review-package adversarial remediations


ROOT = Path(__file__).resolve().parents[1]


def _template_payload() -> dict[str, object]:
    payload = json.loads((ROOT / "docs/templates/deal_template.json").read_text(encoding="utf-8"))
    payload["simulations"] = 1
    return payload


def test_whitespace_only_source_evidence_cannot_verify_a_case(tmp_path):
    payload = _template_payload()
    payload.update(
        {
            "source_status": "VERIFIED",
            "verified_by": "  \t",
            "source_reference": "\n  ",
            "leverage": 0.25,
        }
    )
    input_path = tmp_path / "whitespace-evidence.json"
    input_path.write_text(json.dumps(payload), encoding="utf-8")

    report = (run(input_path, tmp_path / "review") / "investment_committee_report.md").read_text(encoding="utf-8")

    assert "Source status | **REVIEW REQUIRED**" in report
    assert "Decision status | **PASSES INITIAL SCREEN**" not in report


def test_generated_memo_and_report_escape_deal_identifier_html():
    deal = DealInputs(10_000_000, 650_000, leverage=0.25)
    unsafe_identifier = "<script>alert(1)</script>"
    memo = generate_ic_memo(DealCase(unsafe_identifier, "<b>case</b>", deal))
    report = build_screening_package(deal, deal_id=unsafe_identifier, simulations=1)[
        "investment_committee_report.md"
    ].decode("utf-8")

    assert unsafe_identifier not in memo
    assert unsafe_identifier not in report
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in memo
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in report


def test_existing_output_directory_is_rejected_before_overwrite(tmp_path):
    input_path = tmp_path / "deal.json"
    input_path.write_text(json.dumps(_template_payload()), encoding="utf-8")
    output_path = tmp_path / "review"

    run(input_path, output_path)
    original_report = (output_path / "investment_committee_report.md").read_bytes()

    with pytest.raises(ValueError, match="output directory already exists"):
        run(input_path, output_path)

    assert (output_path / "investment_committee_report.md").read_bytes() == original_report


def test_long_hold_workbook_uses_metric_rows_after_cash_flows():
    workbook = load_workbook(
        BytesIO(build_reconciliation_workbook(DealInputs(10_000_000, 650_000, hold_years=15, leverage=0.50))),
        data_only=False,
    )
    cash_flows = workbook["Cash Flows"]
    reconciliation = workbook["Reconciliation"]

    assert cash_flows["A21"].value == "Unlevered IRR"
    assert cash_flows["A24"].value == "Equity multiple"
    assert reconciliation["B8"].value == "='Cash Flows'!B21"
    assert reconciliation["B11"].value == "='Cash Flows'!B24"
    assert all("B11" not in reconciliation.cell(row, 2).value for row in range(8, 12))
