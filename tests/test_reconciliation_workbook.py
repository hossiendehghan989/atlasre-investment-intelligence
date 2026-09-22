from io import BytesIO
from pathlib import Path
from shutil import which
from subprocess import run
from tempfile import TemporaryDirectory

import numpy as np
import numpy_financial as npf
import pytest
from openpyxl import load_workbook

from src.atlasre import DealInputs, underwrite_deal
from src.reconciliation import build_reconciliation_workbook

OFFICE = which("libreoffice") or which("soffice")


def _workbook(deal: DealInputs):
    return load_workbook(BytesIO(build_reconciliation_workbook(deal)), data_only=False)


def test_reconciliation_workbook_has_formula_sheets_and_illustrative_labels():
    workbook = _workbook(DealInputs(10_000_000, 650_000, leverage=0.5))

    assert workbook.sheetnames == [
        "Control",
        "Inputs",
        "Annual NOI",
        "Debt Schedule",
        "Cash Flows",
        "Reconciliation",
        "Lineage",
    ]
    for sheet in workbook.worksheets:
        assert str(sheet["A1"].value).startswith("ILLUSTRATIVE")

    assert workbook["Annual NOI"]["B4"].value == "=Inputs!$B$5*(1+Inputs!$B$7)^(A4-1)"
    assert "IPMT" in workbook["Debt Schedule"]["E4"].value
    assert "PMT" in workbook["Debt Schedule"]["F4"].value
    assert "PPMT" in workbook["Debt Schedule"]["G4"].value
    assert "IRR" in workbook["Cash Flows"]["B11"].value
    assert "NPV" in workbook["Cash Flows"]["B13"].value
    assert workbook["Reconciliation"]["B4"].value == "=Inputs!$B$5/Inputs!$B$4"
    assert workbook["Reconciliation"]["D4"].value == "=B4-C4"


def test_reconciliation_static_outputs_match_independent_python_reevaluation():
    deal = DealInputs(
        purchase_price=8_500_000,
        annual_noi=620_000,
        hold_years=5,
        annual_noi_growth=0.025,
        exit_cap_rate=0.065,
        leverage=0.55,
        debt_rate=0.0625,
        debt_amortization_years=20,
    )
    workbook = _workbook(deal)
    reconciliation = workbook["Reconciliation"]
    outputs = {reconciliation.cell(row, 1).value: reconciliation.cell(row, 3).value for row in range(4, 12)}

    acquisition = deal.purchase_price * (1 + deal.acquisition_cost_pct)
    debt = deal.purchase_price * deal.leverage
    equity = acquisition - debt
    months = deal.hold_years * 12
    monthly_rate = deal.debt_rate / 12
    payment = debt * monthly_rate * (1 + monthly_rate) ** (deal.debt_amortization_years * 12) / (
        (1 + monthly_rate) ** (deal.debt_amortization_years * 12) - 1
    )
    balance = debt
    annual_service = []
    for _year in range(deal.hold_years):
        service = 0.0
        for _month in range(12):
            interest = balance * monthly_rate
            principal = min(max(payment - interest, 0.0), balance)
            service += interest + principal
            balance -= principal
        annual_service.append(service)
    noi = np.asarray([deal.annual_noi * (1 + deal.annual_noi_growth) ** year for year in range(deal.hold_years)])
    exit_value = noi[-1] / deal.exit_cap_rate
    unlevered = np.r_[-acquisition, noi[:-1], noi[-1] + exit_value * (1 - deal.selling_cost_pct)]
    levered = np.r_[-equity, noi[:-1] - np.asarray(annual_service[:-1]), noi[-1] + exit_value * (1 - deal.selling_cost_pct) - annual_service[-1] - balance]

    expected = {
        "Entry cap rate": deal.annual_noi / deal.purchase_price,
        "Exit value": exit_value,
        "Remaining debt at exit": balance,
        "Minimum DSCR": min(noi / np.asarray(annual_service)),
        "Unlevered IRR": float(npf.irr(unlevered)),
        "Levered IRR": float(npf.irr(levered)),
        "Unlevered NPV": sum(flow / (1 + deal.discount_rate) ** period for period, flow in enumerate(unlevered)),
        "Equity multiple": sum(flow for flow in levered if flow > 0) / -sum(flow for flow in levered if flow < 0),
    }
    for metric, value in expected.items():
        np.testing.assert_allclose(outputs[metric], value, rtol=0, atol=1e-6)

    result = underwrite_deal(deal)
    assert outputs["Levered IRR"] == result["levered_irr"]
    np.testing.assert_allclose(outputs["Unlevered NPV"], result["unlevered_npv"], rtol=0, atol=1e-6)
    assert months == 60


@pytest.mark.skipif(OFFICE is None, reason="LibreOffice is not available in this environment")
def test_reconciliation_differences_pass_after_libreoffice_recalculation():
    deal = DealInputs(10_000_000, 650_000, leverage=0.5)
    with TemporaryDirectory() as temporary_directory:
        temporary_path = Path(temporary_directory)
        source = temporary_path / "reconciliation.xlsx"
        recalculated = temporary_path / "recalculated"
        recalculated.mkdir()
        source.write_bytes(build_reconciliation_workbook(deal))
        run(
            [OFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(recalculated), str(source)],
            check=True,
            capture_output=True,
            text=True,
        )
        workbook = load_workbook(recalculated / source.name, data_only=True)

    reconciliation = workbook["Reconciliation"]
    for row in range(4, 12):
        difference = reconciliation.cell(row, 4).value
        tolerance = reconciliation.cell(row, 5).value
        assert reconciliation.cell(row, 6).value == "PASS"
        assert abs(float(difference)) <= float(tolerance)


@pytest.mark.skipif(OFFICE is None, reason="LibreOffice is not available in this environment")
def test_long_hold_workbook_recalculates_after_debt_amortization():
    deal = DealInputs(10_000_000, 650_000, hold_years=25, leverage=0.5, debt_amortization_years=20)
    with TemporaryDirectory() as temporary_directory:
        temporary_path = Path(temporary_directory)
        source = temporary_path / "long-hold.xlsx"
        recalculated = temporary_path / "recalculated"
        recalculated.mkdir()
        source.write_bytes(build_reconciliation_workbook(deal))
        run(
            [OFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(recalculated), str(source)],
            check=True,
            capture_output=True,
            text=True,
        )
        workbook = load_workbook(recalculated / source.name, data_only=True)

    debt_schedule = workbook["Debt Schedule"]
    first_post_amortization_row = 4 + 20 * 12
    assert debt_schedule.cell(first_post_amortization_row, 5).value == 0
    assert debt_schedule.cell(first_post_amortization_row, 6).value == 0
    assert debt_schedule.cell(first_post_amortization_row, 7).value == 0
    assert all(workbook["Reconciliation"].cell(row, 6).value == "PASS" for row in range(4, 12))
